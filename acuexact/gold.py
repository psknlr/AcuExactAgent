"""Explicit-column gold import; no inferred article identity or silent quote splitting."""
import csv
from pathlib import Path


def import_gold(path, article_column="文章ID", quote_column="古籍条文", book_column=None, chapter_column=None, sheet=None):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        ws = workbook[sheet] if sheet else workbook.active
        values = list(ws.iter_rows(values_only=True))
        workbook.close()
        if not values:
            raise ValueError("Gold workbook is empty.")
        headers = [str(x or "") for x in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:]]
    else:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream, delimiter="\t" if path.suffix.lower() == ".tsv" else ",")
            headers, rows = reader.fieldnames or [], list(reader)
    required = [article_column, quote_column] + [x for x in (book_column, chapter_column) if x]
    missing = set(required) - set(headers)
    if missing:
        raise ValueError(f"Missing gold columns: {sorted(missing)}")
    result = []
    for index, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row.values()):
            continue
        if row.get(article_column) in (None, "") or row.get(quote_column) in (None, ""):
            raise ValueError(f"Gold row {index} lacks an article ID or quote; refusing to guess.")
        q = {"id": f"gold:row{index}", "article_id": str(row[article_column]), "text": str(row[quote_column]), "gold_row": index}
        if book_column:
            q["source_explicit"] = {"book": str(row[book_column]) if row[book_column] else None}
            if chapter_column:
                q["source_explicit"]["chapter"] = str(row[chapter_column]) if row.get(chapter_column) else None
        result.append(q)
    return result
